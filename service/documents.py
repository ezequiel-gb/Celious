"""Bounded DOCX/ODT creation and selective text edits; no layout rendering or execution."""
import io
import json
import re
import zipfile
from lxml import etree as XML

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
T = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
O = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
NS = dict(text=T, office=O, table='urn:oasis:names:tc:opendocument:xmlns:table:1.0',
          style='urn:oasis:names:tc:opendocument:xmlns:style:1.0',
          fo='urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0')
MIME = 'application/vnd.oasis.opendocument.text'


def parse(data):
    if re.search(br'<!\s*(DOCTYPE|ENTITY)', data, re.I):
        raise ValueError('Document entities and DTDs are unsupported')
    try:
        tree = XML.fromstring(data, XML.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))
        if tree.getroottree().docinfo.doctype:
            raise ValueError('Document entities and DTDs are unsupported')
        return tree
    except XML.XMLSyntaxError:
        raise ValueError('Invalid document XML') from None


def package(data, editing=False):
    try:
        bundle = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ValueError('Invalid document ZIP') from None
    entries = bundle.infolist()
    names = [e.filename for e in entries]
    if (len(entries) > 1000 or len(names) != len(set(names)) or sum(e.file_size for e in entries) > 50*1024*1024
            or any(e.flag_bits & 1 or e.filename.startswith(('/', '\\')) or '..' in e.filename.replace('\\', '/').split('/') for e in entries)):
        bundle.close()
        raise ValueError('Encrypted, duplicated, oversized or unsafe document package')
    if editing and any('vbaproject' in n.lower() or '_xmlsignatures/' in n.lower() or 'documentsignatures.xml' in n.lower() or n.lower().startswith(('basic/', 'scripts/')) for n in names):
        bundle.close()
        raise ValueError('Signed or macro-bearing documents cannot be edited safely')
    return bundle


def parts(bundle, suffix, reading=False):
    if suffix == '.docx':
        if 'word/document.xml' not in bundle.namelist():
            raise ValueError('Missing DOCX document part')
        return [n for n in bundle.namelist() if n == 'word/document.xml' or re.fullmatch(r'word/(header|footer)\d+\.xml', n)]
    if suffix == '.odt':
        if reading and 'content.xml' in bundle.namelist() and 'mimetype' not in bundle.namelist():
            return ['content.xml']  # Preserve legacy extraction-only fixtures; edits require a full package.
        if not {'mimetype', 'META-INF/manifest.xml', 'content.xml'} <= set(bundle.namelist()):
            raise ValueError('Missing ODT package parts')
        manifest = parse(bundle.read('META-INF/manifest.xml'))
        if bundle.read('mimetype').decode() != MIME or any(isinstance(n.tag, str) and n.tag.rsplit('}',1)[-1] == 'encryption-data' for n in manifest.iter()):
            raise ValueError('Unsupported or encrypted ODT')
        return ['content.xml']
    raise ValueError('Document recipes and formatted edits support DOCX and ODT')


def paragraphs(tree, suffix):
    return list(tree.iter('{'+W+'}p')) if suffix == '.docx' else [n for n in tree.iter() if n.tag in {'{'+T+'}p', '{'+T+'}h'}]


def tokens(paragraph, suffix):
    result = []
    if suffix == '.docx':
        for node in paragraph.iter():
            # Do not count text in nested text-box paragraphs twice.
            if next((p for p in node.iterancestors() if p.tag == '{'+W+'}p'), None) is not paragraph:
                continue
            if node.tag == '{'+W+'}t': result.append((node, 'text', node.text or ''))
            elif node.tag in {'{'+W+'}tab', '{'+W+'}br', '{'+W+'}cr'}:
                result.append((node, None, '\t' if node.tag.endswith('tab') else '\n'))
    else:
        def visit(node):
            if node.text: result.append((node, 'text', node.text))
            for child in node:
                if not isinstance(child.tag, str):
                    pass  # XML comments/processing instructions are not visible paragraph text.
                elif child.tag == '{'+T+'}s':
                    count = int(child.get('{'+T+'}c', '1'))
                    if not 1 <= count <= 10000: raise ValueError('Invalid ODT space repetition')
                    result.append((child, None, ' '*count))
                elif child.tag in {'{'+T+'}tab', '{'+T+'}line-break'}:
                    result.append((child, None, '\t' if child.tag.endswith('tab') else '\n'))
                else: visit(child)
                if child.tail: result.append((child, 'tail', child.tail))
        visit(paragraph)
    return result


def extract(data, suffix):
    with package(data) as bundle:
        return '\n'.join(''.join(t[2] for t in tokens(p, suffix)) for name in parts(bundle, suffix, reading=True)
                         for p in paragraphs(parse(bundle.read(name)), suffix))


def edit(data, suffix, old, new):
    if not isinstance(old, str) or not old or not isinstance(new, str) or any(c in new for c in '\r\n\t'):
        raise ValueError('Formatted editing requires unique text within one paragraph and a single-line replacement')
    if suffix == '.odt' and ('  ' in new or new.startswith(' ') or new.endswith(' ')):
        raise ValueError('ODT replacement uses ordinary single spaces; explicit repeated-space editing is unsupported')
    with package(data, editing=True) as bundle:
        trees = {name: parse(bundle.read(name)) for name in parts(bundle, suffix)}
        if suffix == '.odt' and any(n.tag == '{'+O+'}scripts' for tree in trees.values() for n in tree.iter()):
            raise ValueError('Script-bearing documents cannot be edited safely')
        found = []
        for name, tree in trees.items():
            for paragraph in paragraphs(tree, suffix):
                pieces = tokens(paragraph, suffix)
                text = ''.join(t[2] for t in pieces)
                for match in re.finditer(re.escape(old), text): found.append((name, paragraph, pieces, match.start(), match.end()))
        if len(found) != 1:
            raise ValueError('old_text must occur exactly once in supported document paragraphs')
        name, paragraph, pieces, start, end = found[0]
        if suffix == '.docx' and any(n.tag in {'{'+W+'}fldChar', '{'+W+'}instrText', '{'+W+'}ins', '{'+W+'}del'} for n in [*paragraph.iter(), *paragraph.iterancestors()]):
            raise ValueError('Editing paragraphs with fields or tracked changes is unsupported')
        if suffix == '.odt' and any(isinstance(n.tag, str) and n.tag.startswith('{'+T+'}') and n.tag.split('}', 1)[1] not in
                {'p', 'h', 'span', 'a', 's', 'tab', 'line-break', 'bookmark', 'bookmark-start', 'bookmark-end'} for n in paragraph.iter()):
            raise ValueError('Editing ODT fields or change-tracking content is unsupported')
        offset, selected = 0, []
        for node, attr, text in pieces:
            lo, hi = max(0, start-offset), min(len(text), end-offset)
            if lo < hi:
                if attr is None: raise ValueError('Editing across explicit whitespace elements is unsupported')
                selected.append((node, attr, text, lo, hi))
            offset += len(text)
        for i, (node, attr, text, lo, hi) in enumerate(selected):
            setattr(node, attr, text[:lo] + (new if i == 0 else '') + text[hi:])
            if suffix == '.docx': node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        changed = XML.tostring(trees[name], xml_declaration=True, encoding='UTF-8', standalone=True)
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as target:
            for entry in bundle.infolist(): target.writestr(entry, changed if entry.filename == name else bundle.read(entry))
        result = output.getvalue()
        if len(result) > 10*1024*1024: raise ValueError('Edited document exceeds 10 MB')
        return result


def recipe(content):
    if len(content.encode()) > 1000000: raise ValueError('Document recipe exceeds 1 MB')
    value = json.loads(content)
    if not isinstance(value, dict) or set(value) - {'title', 'blocks'} or not isinstance(value.get('blocks'), list) or not 1 <= len(value['blocks']) <= 1000:
        raise ValueError('Document recipe requires 1–1000 blocks and optional title')
    if not isinstance(value.get('title', ''), str): raise ValueError('Document title must be text')
    cells = 0
    for block in value['blocks']:
        if not isinstance(block, dict) or block.get('type') not in {'paragraph', 'heading', 'table'}:
            raise ValueError('Document blocks: paragraph, heading or table')
        if block['type'] == 'table':
            rows = block.get('rows')
            if set(block) != {'type', 'rows'} or not isinstance(rows, list) or not rows or not isinstance(rows[0], list) or not rows[0]: raise ValueError('Table requires rectangular text rows')
            for row in rows:
                if not isinstance(row, list) or len(row) != len(rows[0]) or not all(isinstance(v, str) for v in row): raise ValueError('Table requires rectangular text rows')
                cells += len(row)
        elif (set(block) - {'type', 'text', 'level'} or not isinstance(block.get('text'), str)
              or block['type'] == 'heading' and (type(block.get('level', 1)) is not int or not 1 <= block.get('level', 1) <= 6)):
            raise ValueError('Paragraph/heading requires text; heading level is 1–6')
    if cells > 10000: raise ValueError('Document tables exceed 10000 cells')
    return value


def create(content, suffix):
    value = recipe(content)
    out = io.BytesIO()
    if suffix == '.docx':
        from docx import Document
        doc = Document()
        doc.core_properties.title = value.get('title', '')
        for block in value['blocks']:
            if block['type'] == 'table':
                table = doc.add_table(rows=len(block['rows']), cols=len(block['rows'][0]))
                for row, values in zip(table.rows, block['rows']):
                    for cell, text in zip(row.cells, values): cell.text = text
            elif block['type'] == 'heading': doc.add_heading(block['text'], block.get('level', 1))
            else: doc.add_paragraph(block['text'])
        doc.save(out)
    elif suffix == '.odt':
        root = XML.Element('{'+O+'}document-content', nsmap=NS, attrib={'{'+O+'}version': '1.2'})
        body = XML.SubElement(XML.SubElement(root, '{'+O+'}body'), '{'+O+'}text')
        for i, block in enumerate(value['blocks']):
            if block['type'] == 'table':
                table = XML.SubElement(body, '{'+NS['table']+'}table', attrib={'{'+NS['table']+'}name': 'Table'+str(i)})
                for row in block['rows']:
                    target = XML.SubElement(table, '{'+NS['table']+'}table-row')
                    for text in row:
                        cell = XML.SubElement(target, '{'+NS['table']+'}table-cell', attrib={'{'+O+'}value-type': 'string'})
                        odt_text(XML.SubElement(cell, '{'+T+'}p'), text)
            else:
                node = XML.SubElement(body, '{'+T+'}h' if block['type'] == 'heading' else '{'+T+'}p')
                if block['type'] == 'heading': node.set('{'+T+'}outline-level', str(block.get('level', 1)))
                odt_text(node, block['text'])
        manifest_ns = 'urn:oasis:names:tc:opendocument:xmlns:manifest:1.0'
        manifest = XML.Element('{'+manifest_ns+'}manifest', nsmap={'manifest': manifest_ns}, attrib={'{'+manifest_ns+'}version': '1.2'})
        meta = XML.Element('{'+O+'}document-meta', nsmap={'office': O, 'dc': 'http://purl.org/dc/elements/1.1/'}, attrib={'{'+O+'}version': '1.2'})
        XML.SubElement(XML.SubElement(meta, '{'+O+'}meta'), '{http://purl.org/dc/elements/1.1/}title').text = value.get('title', '')
        for name, media in [('/', MIME), ('content.xml', 'text/xml'), ('meta.xml', 'text/xml')]:
            XML.SubElement(manifest, '{'+manifest_ns+'}file-entry', attrib={'{'+manifest_ns+'}full-path': name, '{'+manifest_ns+'}media-type': media})
        with zipfile.ZipFile(out, 'w') as bundle:
            bundle.writestr('mimetype', MIME, compress_type=zipfile.ZIP_STORED)
            bundle.writestr('content.xml', XML.tostring(root, xml_declaration=True, encoding='UTF-8'), compress_type=zipfile.ZIP_DEFLATED)
            bundle.writestr('meta.xml', XML.tostring(meta, xml_declaration=True, encoding='UTF-8'), compress_type=zipfile.ZIP_DEFLATED)
            bundle.writestr('META-INF/manifest.xml', XML.tostring(manifest, xml_declaration=True, encoding='UTF-8'), compress_type=zipfile.ZIP_DEFLATED)
    else: raise ValueError('Document creation supports DOCX and ODT')
    if len(out.getvalue()) > 10*1024*1024: raise ValueError('Generated document exceeds 10 MB')
    return out.getvalue()


def odt_text(node, text):
    """Encode non-collapsible whitespace explicitly in generated ODF paragraphs."""
    position, last = 0, None
    for match in re.finditer(r' {2,}|^ | $|\t|\r\n|\r|\n', text):
        if last is None: node.text = (node.text or '') + text[position:match.start()]
        else: last.tail = (last.tail or '') + text[position:match.start()]
        value = match.group()
        tag = 's' if value.startswith(' ') else 'tab' if value == '\t' else 'line-break'
        last = XML.SubElement(node, '{'+T+'}'+tag)
        if tag == 's': last.set('{'+T+'}c', str(len(value)))
        position = match.end()
    if last is None: node.text = text
    else: last.tail = (last.tail or '') + text[position:]
