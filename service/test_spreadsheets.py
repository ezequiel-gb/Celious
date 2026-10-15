import asyncio
import base64
from datetime import datetime
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from engine import Files
from spreadsheets import Spreadsheets, SpreadsheetError, load, inspect, read, search, encode, sheet, put, scalar


def ods(path, tables):
    content=('<?xml version="1.0"?><office:document-content '
             'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
             'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
             'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
             '<office:body><office:spreadsheet>'+tables+'</office:spreadsheet></office:body></office:document-content>')
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet')
        z.writestr('content.xml',content)


class SpreadsheetTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.base=Path(self.tmp.name)
        self.root=self.base/'data'; self.root.mkdir()
        self.f=Files({'direct_operations':True,'roots':{'Test':{'path':str(self.root),'writable':True}}},self.base/'runtime')
        self.s=Spreadsheets(self.f)
        self.spec=[{'name':'Español','rows':[['Nombre','Valor','Activo'],['José',12,True],['José',None,False],['Ana',3.5,None]]},
                   {'name':'Other','rows':[['Date'],['2026-10-14']]}]

    def tearDown(self):
        self.tmp.cleanup()

    def create(self, name='book.xlsx'):
        return self.s.write('Test',name,json.dumps(self.spec))

    def test_xlsx_types_unicode_multisheet(self):
        self.create()
        b=self.s.load('Test','book.xlsx')
        self.assertEqual([s['name'] for s in b],['Español','Other'])
        page=read(b,'Español',2,1,1,3)
        self.assertEqual([v['type'] for v in page['rows'][0]['cells']],['string','number','boolean'])
        self.assertEqual(page['rows'][0]['cells'][0]['value'],'José')
        self.assertEqual(page['next_row'],3)
        self.assertEqual(read(b,'Español',3,1)['rows'][0]['cells'][1]['type'],'empty')
        with self.assertRaises(SpreadsheetError): read(b)

    def test_ods_typed_formulas_repeats_merges(self):
        ods(self.root/'typed.ods','<table:table table:name="Datos"><table:table-row>'
            '<table:table-cell office:value-type="string"><text:p>España<text:s text:c="2"/>sí</text:p></table:table-cell>'
            '<table:table-cell office:value-type="float" office:value="2.5"/>'
            '<table:table-cell office:value-type="boolean" office:boolean-value="true"/>'
            '<table:table-cell office:value-type="date" office:date-value="2026-10-14"/>'
            '<table:table-cell office:value-type="float" office:value="3" table:formula="of:=1+2"/>'
            '<table:table-cell table:number-columns-spanned="2" office:value-type="string"><text:p>Merge</text:p></table:table-cell>'
            '<table:covered-table-cell/></table:table-row>'
            '<table:table-row table:number-rows-repeated="3"><table:table-cell office:value-type="string"><text:p>Repeated</text:p></table:table-cell></table:table-row>'
            '<table:table-row table:number-rows-repeated="10000"><table:table-cell table:number-columns-repeated="1000"/></table:table-row></table:table>'
            '<table:table table:name="Second"><table:table-row/></table:table>')
        b=load(self.root/'typed.ods'); s=b[0]
        self.assertEqual(s['cells'][1,1]['value'],'España  sí')
        self.assertEqual(s['cells'][1,2]['type'],'number')
        self.assertEqual(s['cells'][1,3]['value'],True)
        self.assertEqual(s['cells'][1,4]['type'],'date')
        self.assertEqual(s['cells'][1,5]['formula'],'of:=1+2')
        self.assertEqual(s['cells'][1,5]['cached'],3)
        self.assertEqual(s['merged_ranges'],['F1:G1'])
        self.assertEqual(s['rows'],4)
        self.assertIn('of:=1+2',self.f.read('Test','typed.ods')['text'])

    def test_formula_cache_xlsx(self):
        from openpyxl import Workbook
        w=Workbook(); w.active['A1']='=1+2'; w.save(self.root/'f.xlsx'); w.close()
        cell=load(self.root/'f.xlsx')[0]['cells'][1,1]
        self.assertEqual(cell['formula'],'=1+2'); self.assertIsNone(cell['cached'])
        self.assertIn('=1+2',self.f.read('Test','f.xlsx')['text'])

    def test_native_dates(self):
        from openpyxl import Workbook
        w=Workbook(); w.active['A1']=datetime(2026,10,14); w.save(self.root/'d.xlsx'); w.close()
        self.assertEqual(load(self.root/'d.xlsx')[0]['cells'][1,1]['type'],'date')

    def test_csv_tsv_utf8_quotes_empty_literal(self):
        for ext,sep in [('.csv',','),('.tsv','\t')]:
            (self.root/('a'+ext)).write_text(f'Name{sep}Value\n"Gómez{sep}José"{sep}\n=1+2{sep}003\n',encoding='utf-8-sig')
            b=load(self.root/('a'+ext))
            self.assertEqual(b[0]['cells'][3,2]['value'],'003')
            self.assertEqual(b[0]['cells'][3,1]['type'],'string')
            with self.assertRaises(SpreadsheetError): encode(b,ext)
            self.s.convert('Test','a'+ext,'converted'+ext+'.xlsx',acknowledge_loss=True)
            self.assertNotIn('formula',load(self.root/('converted'+ext+'.xlsx'))[0]['cells'][3,1])

    def test_duplicate_missing_inconsistent_search(self):
        self.create(); b=load(self.root/'book.xlsx')
        self.assertEqual(search(b,'Español',columns=[1],mode='duplicates')['matches'],[{'rows':[2,3]}])
        self.assertEqual(search(b,'Español',columns=[2],mode='missing')['matches'],[{'row':3,'columns':[2]}])
        self.assertEqual(search(b,'Español',query='JOSÉ',columns=[1],limit=1)['next_offset'],1)
        self.assertEqual(search(b,'Español',columns=[2],mode='inconsistent')['total'],0)

    def test_pagination_large_sparse_dimensions_blank_counts(self):
        s=sheet('Large')
        for r in range(1,1501): put(s,r,1,scalar(r))
        put(s,2000,3,scalar('last'))
        data=[s]
        self.assertEqual(read(data,row_count=500,column_count=3)['next_row'],501)
        self.assertEqual(read(data,start_row=1501,row_count=499,column_count=3)['rows'][0]['cells'][0]['type'],'empty')
        self.assertEqual(inspect(data)['worksheets'][0]['blank_rows_within_extent'],499)
        self.assertEqual(inspect(data)['worksheets'][0]['blank_columns_within_extent'],1)
        for kwargs in [{'row_count':501},{'column_count':101},{'start_row':0},{'row_count':500,'column_count':100}]:
            with self.assertRaises(SpreadsheetError): read(data,**kwargs)

    def test_conversion_roundtrip_and_selection(self):
        self.create('source.ods')
        before=(self.root/'source.ods').read_bytes()
        self.s.convert('Test','source.ods','result.xlsx',acknowledge_loss=True)
        original,converted=load(self.root/'source.ods'),load(self.root/'result.xlsx')
        self.assertEqual(original,converted)
        self.s.convert('Test','result.xlsx','roundtrip.ods',acknowledge_loss=True)
        self.assertEqual(original,load(self.root/'roundtrip.ods'))
        with self.assertRaises(SpreadsheetError): self.s.convert('Test','source.ods','bad.csv',acknowledge_loss=True)
        self.s.convert('Test','source.ods','selected.csv',worksheet='Other',acknowledge_loss=True)
        self.assertEqual((self.root/'source.ods').read_bytes(),before)

    def test_loss_and_formula_rejection(self):
        self.create()
        with self.assertRaises(SpreadsheetError): self.s.convert('Test','book.xlsx','out.ods')
        with self.assertRaises(SpreadsheetError): self.s.edit('Test','book.xlsx','out.xlsx','[]')
        from openpyxl import Workbook
        w=Workbook(); w.active['A1']='=1+1'; w.save(self.root/'f.xlsx'); w.close()
        with self.assertRaises(SpreadsheetError): self.s.convert('Test','f.xlsx','out.ods',acknowledge_loss=True)
        self.s.convert('Test','f.xlsx','out.ods',acknowledge_loss=True,formula_mode='cached_values')

    def test_edit_preserves_common_features_and_literal_text(self):
        self.create()
        from openpyxl import load_workbook
        from openpyxl.styles import Font
        w=load_workbook(self.root/'book.xlsx'); w['Español']['A2'].font=Font(bold=True); w['Español']['D2']='=B2*2'; w.save(self.root/'book.xlsx'); w.close()
        ops=[{'action':'set_cell','sheet':'Español','row':2,'column':1,'value':'=untrusted'},
             {'action':'append_rows','sheet':'Español','rows':[['Nuevo',5,True]]},
             {'action':'add_sheet','sheet':'New'}]
        self.s.edit('Test','book.xlsx','edited.xlsx',json.dumps(ops),True)
        w=load_workbook(self.root/'edited.xlsx')
        self.assertTrue(w['Español']['A2'].font.bold)
        self.assertEqual(w['Español']['A2'].data_type,'s')
        self.assertEqual(w['Español']['D2'].value,'=B2*2')
        self.assertEqual(w['Español']['A5'].value,'Nuevo'); self.assertIn('New',w.sheetnames); w.close()

    def test_controlled_formula_write(self):
        self.create()
        op={'action':'set_cell','sheet':'Español','row':2,'column':4,'formula':'=SUM(B2:B4)'}
        with self.assertRaises(SpreadsheetError): self.s.edit('Test','book.xlsx','out.xlsx',json.dumps([op]),True)
        self.s.edit('Test','book.xlsx','out.xlsx',json.dumps([op]),True,True)
        self.assertEqual(load(self.root/'out.xlsx')[0]['cells'][2,4]['formula'],'=SUM(B2:B4)')
        for formula in ['=WEBSERVICE("url")','=HYPERLINK("x")','=UNKNOWN(1)','=Sheet1!A1+[book]A1']:
            op['formula']=formula
            with self.assertRaises(SpreadsheetError): self.s.edit('Test','book.xlsx','bad.xlsx',json.dumps([op]),True,True)

    def test_backup_conflict_atomic_failure_recovery(self):
        result=self.create(); original=(self.root/'book.xlsx').read_bytes()
        result=self.s.write('Test','book.xlsx',json.dumps([{'name':'New','rows':[[1]]}]),expected_sha256=result['sha256'])
        backup=self.f.runtime/'backups'/result['backup_id']/'original.bin'
        self.assertEqual(backup.read_bytes(),original)
        current=(self.root/'book.xlsx').read_bytes()
        with self.assertRaises(ValueError): self.create()
        with patch('engine.os.replace',side_effect=OSError('simulated disk failure')):
            with self.assertRaises(OSError): self.s.write('Test','book.xlsx',json.dumps(self.spec),expected_sha256=result['sha256'])
        self.assertEqual((self.root/'book.xlsx').read_bytes(),current)
        self.assertFalse(list(self.root.glob('celious-*')))
        self.s.save('Test','restored.xlsx',backup.read_bytes())
        self.assertEqual((self.root/'restored.xlsx').read_bytes(),original)

    def test_failed_conversion_preserves_source_and_destination(self):
        self.create(); before=(self.root/'book.xlsx').read_bytes()
        (self.root/'existing.ods').write_bytes(b'existing')
        with self.assertRaises(ValueError): self.s.convert('Test','book.xlsx','existing.ods',acknowledge_loss=True)
        self.assertEqual((self.root/'existing.ods').read_bytes(),b'existing')
        self.assertEqual((self.root/'book.xlsx').read_bytes(),before)
        with patch('spreadsheets.load',side_effect=SpreadsheetError('verify failed')):
            with self.assertRaises(SpreadsheetError): self.s.save('Test','bad.xlsx',before)
        self.assertFalse((self.root/'bad.xlsx').exists())

    def test_security_roots_direct_disabled_readonly_links(self):
        self.create()
        for path in ['../book.xlsx','C:\\Windows\\a.xlsx','.hidden.xlsx','CON.xlsx']:
            with self.assertRaises(PermissionError): self.s.load('Test',path)
            with self.assertRaises(PermissionError): self.s.write('Test',path,json.dumps(self.spec))
        self.f.config['read_only']=True
        with self.assertRaises(PermissionError): self.s.write('Test','new.xlsx',json.dumps(self.spec))
        self.f.config['read_only']=False; self.f.config['direct_operations']=False
        with self.assertRaises(PermissionError): self.s.write('Test','new.xlsx',json.dumps(self.spec))

    def test_malformed_mismatched_entities_repetitions(self):
        for ext in ['ods','xlsx','xls','csv']:
            (self.root/('bad.'+ext)).write_bytes(b'\x00not a document')
            with self.assertRaises(SpreadsheetError): load(self.root/('bad.'+ext))
        ods(self.root/'repeat.ods','<table:table table:name="X"><table:table-row table:number-rows-repeated="999999999"/></table:table>')
        with self.assertRaises(SpreadsheetError): load(self.root/'repeat.ods')
        with zipfile.ZipFile(self.root/'entity.ods','w') as z:
            z.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet')
            z.writestr('content.xml','<!DOCTYPE x [<!ENTITY x "evil">]><x>&x;</x>')
        with self.assertRaises(SpreadsheetError): load(self.root/'entity.ods')

    def test_verified_output_and_audit_privacy(self):
        result=self.create()
        self.assertTrue(result['verified'])
        self.assertNotIn('José',(self.f.runtime/'audit.jsonl').read_text())

    def test_search_zero_and_false_are_not_empty(self):
        s=sheet('Values'); put(s,1,1,scalar('Header')); put(s,2,1,scalar(0)); put(s,3,1,scalar(False))
        self.assertEqual(search([s],query='0')['matches'],[{'row':2,'columns':[1]}])
        self.assertEqual(search([s],query='false')['matches'],[{'row':3,'columns':[1]}])
        self.assertEqual(search([s],mode='missing')['total'],0)

    def test_no_silent_string_truncation_or_invalid_names(self):
        with self.assertRaises(SpreadsheetError): self.s.write('Test','long.xlsx',json.dumps([{'name':'X','rows':[['a'*32768]]}]))
        self.assertFalse((self.root/'long.xlsx').exists())
        for name in ['',1,'bad/name','X'*32]:
            with self.assertRaises(SpreadsheetError): self.s.write('Test','bad.xlsx',json.dumps([{'name':name,'rows':[[1]]}]))

    def test_link_policy_without_os_fixture_creation(self):
        from types import SimpleNamespace
        import stat
        fake=SimpleNamespace(st_mode=stat.S_IFREG,st_nlink=2,st_file_attributes=0)
        with patch('pathlib.Path.lstat',return_value=fake):
            with self.assertRaises(PermissionError): self.f._no_links(self.root/'fixture.xlsx')


class ProtocolTests(unittest.TestCase):
    def test_registered_tools_and_real_calls(self):
        import shutil
        temporary=tempfile.TemporaryDirectory(dir=Path(__file__).parent)
        self.addCleanup(temporary.cleanup)
        base=Path(temporary.name)
        service=base/'service'; service.mkdir()
        for name in ['server.py','engine.py','spreadsheets.py','memory_tools.py','project_memory.py','second_brain.py',
                     'operations_state.py','ingestion.py','workflows.py','workspace_tools.py']:
            shutil.copy2(Path(__file__).parent/name,service/name)
        (base/'vault/Inbox').mkdir(parents=True)
        (service/'config.json').write_text(json.dumps({'direct_operations':True,'roots':{'Vault':{'path':str(base/'vault'),'writable':True}}}),'utf-8')
        async def run():
            from mcp import Client
            from mcp.client.stdio import StdioServerParameters
            import uuid
            async with Client(StdioServerParameters(command=sys.executable,args=[str(service/'server.py')])) as client:
                names={t.name for t in (await client.list_tools()).tools}
                self.assertTrue({'inspect_spreadsheet','read_spreadsheet','search_spreadsheet','convert_spreadsheet','edit_spreadsheet','write_spreadsheet'}<=names)
                # No production data: use a copied server and disposable configured root.
                path='Inbox/spreadsheet-test-'+uuid.uuid4().hex+'.xlsx'
                r=await client.call_tool('write_spreadsheet',{'root':'Vault','destination':path,'workbook_json':json.dumps([{'name':'Datos','rows':[['Nombre'],['José']]}])})
                self.assertFalse(r.is_error,r)
                try:
                    for tool,args in [('inspect_spreadsheet',{}),('read_spreadsheet',{'row_count':2}),('search_spreadsheet',{'query':'José'})]:
                        r=await client.call_tool(tool,{'root':'Vault','path':path,**args})
                        self.assertFalse(r.is_error,r)
                    converted=path+'.ods'
                    r=await client.call_tool('convert_spreadsheet',{'root':'Vault','path':path,'destination':converted,'acknowledge_loss':True})
                    self.assertFalse(r.is_error,r)
                    r=await client.call_tool('edit_spreadsheet',{'root':'Vault','path':path,'destination':path+'.xlsx','acknowledge_loss':True,
                        'operations_json':json.dumps([{'action':'set_cell','sheet':'Datos','row':2,'column':1,'value':'Ana'}])})
                    self.assertFalse(r.is_error,r)
                    r=await client.call_tool('read_spreadsheet',{'root':'Vault','path':'../blocked.xlsx'})
                    self.assertTrue(r.is_error)
                finally:
                    r=await client.call_tool('delete_file',{'root':'Vault','path':path})
                    self.assertFalse(r.is_error,r)
        asyncio.run(asyncio.wait_for(run(),45))


if __name__=='__main__':
    unittest.main(verbosity=2)
