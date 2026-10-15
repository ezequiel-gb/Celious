"""Celious's local filesystem policy. No network or shell execution."""
from pathlib import Path, PureWindowsPath
import difflib
import base64
import hashlib
import json
import os
import re
import stat
import tempfile
import threading
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone

BASE = Path(__file__).resolve().parent
TEXT = {'.md', '.txt', '.csv', '.tsv', '.json', '.yaml', '.yml', '.xml', '.html', '.css', '.py', '.js', '.ts', '.toml', ''}
WRITE = {'.md', '.txt', '.csv', '.tsv', '.json'}
LIMIT = 10 * 1024 * 1024


class Files:
    def __init__(self, config=None, runtime=None):
        self.config = config or json.loads((BASE / 'config.json').read_text('utf-8-sig'))
        self.runtime = Path(runtime or BASE / '.runtime')
        self.runtime.mkdir(exist_ok=True, parents=True)
        self.lock = threading.RLock()
        self.roots = {k: {**v, 'path': Path(v['path']).absolute()} for k, v in self.config['roots'].items()}
        for root in self.roots.values():
            if not root['path'].is_dir():
                raise ValueError('Configured root does not exist')
            self._no_links(root['path'])

    def audit(self, action, root='', path='', status='ok', request_id=''):
        # Do not record content, credentials, or raw rejected absolute paths.
        event = dict(time=datetime.now(timezone.utc).isoformat(), action=action, root=root,
                     path=path if len(path) < 300 else '[long path]', status=status, request_id=request_id)
        with self.lock, (self.runtime / 'audit.jsonl').open('a', encoding='utf-8') as out:
            out.write(json.dumps(event) + '\n')

    def _no_links(self, p):
        for node in [p, *p.parents]:
            try:
                s = node.lstat()
            except FileNotFoundError:
                continue
            if stat.S_ISLNK(s.st_mode) or getattr(s, 'st_file_attributes', 0) & 0x400:
                raise PermissionError('Symbolic links and Windows reparse points are blocked')
            if stat.S_ISREG(s.st_mode) and s.st_nlink > 1:
                raise PermissionError('Hard-linked files are blocked')

    def path(self, root, relative='', write=False, _aliases=True):
        try:
            if root not in self.roots:
                raise PermissionError('Unknown root')
            if not isinstance(relative, str) or '\x00' in relative or ':' in relative:
                raise PermissionError('Invalid relative path')
            win = PureWindowsPath(relative)
            parts = relative.replace('\\', '/').split('/')
            if win.is_absolute() or win.drive or relative.startswith(('/', '\\')):
                raise PermissionError('Absolute paths are blocked; use a root alias')
            if any(x == '..' or x.startswith('.') or x.endswith((' ', '.')) for x in parts if x):
                raise PermissionError('Traversal, hidden paths and ambiguous names are blocked')
            reserved = {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}
            if any(x.split('.')[0].upper() in reserved or any(c in x for c in '<>"|?*') for x in parts):
                raise PermissionError('Windows special paths are blocked')
            normalized_parts=[x for x in parts if x]
            aliases=self.config.get('read_path_aliases',{}).get(root,{}) if _aliases else {}
            for old,target in sorted(aliases.items(),key=lambda item:len(item[0]),reverse=True):
                old_parts=old.replace('\\', '/').split('/')
                if [x.casefold() for x in normalized_parts[:len(old_parts)]] == [x.casefold() for x in old_parts]:
                    if write:
                        raise PermissionError('Historical path is read-only; use the canonical archive path')
                    resolved='/'.join([target,*normalized_parts[len(old_parts):]])
                    # Validate alias targets under exactly the same root/policy, without chains.
                    return self.path(root,resolved,_aliases=False)
            spec = self.roots[root]
            if write and (self.config.get('read_only', False) or not spec.get('writable', False)):
                raise PermissionError('This root is read-only')
            p = spec['path'].joinpath(*[x for x in parts if x])
            if not p.resolve().is_relative_to(spec['path'].resolve()):
                raise PermissionError('Outside authorized root')
            self._no_links(p)
            return p
        except (PermissionError, ValueError):
            self.audit('access', status='denied')
            raise

    def roots_info(self):
        direct = self.config.get('direct_operations', False)
        return {k: {'path': str(v['path']), 'writable': v.get('writable', False) and not self.config.get('read_only', False),
                    'direct_operations': direct, 'deletion': 'recoverable' if direct else 'unavailable'} for k, v in self.roots.items()}

    def mutation_path(self, root, relative):
        if not self.config.get('direct_operations', False):
            raise PermissionError('Direct operations are disabled in local configuration')
        p = self.path(root, relative, write=True)
        # A root or ancestor of another configured root cannot be changed.
        if any(v['path'].resolve().is_relative_to(p.resolve()) for v in self.roots.values()):
            raise PermissionError('Cannot modify an authorized root or its ancestor')
        return p

    def _backup(self, p, root, relative, action):
        ident = uuid.uuid4().hex
        folder = self.runtime / 'backups' / ident
        folder.mkdir(parents=True)
        data = p.read_bytes()
        saved = folder / 'original.bin'
        saved.write_bytes(data)
        if saved.read_bytes() != data:
            raise OSError('Backup integrity verification failed')
        (folder / 'record.json').write_text(json.dumps({'root': root, 'path': relative, 'action': action}), 'utf-8')
        return ident

    def raw_read(self, root, relative, offset=0, max_bytes=65536):
        p = self.path(root, relative)
        if not p.is_file():
            raise ValueError('Not a file')
        offset, max_bytes = max(0, offset), min(262144, max(1, max_bytes))
        with p.open('rb') as src:
            src.seek(offset)
            data = src.read(max_bytes)
        size = p.stat().st_size
        self.audit('read_binary', root, relative)
        return {'base64': base64.b64encode(data).decode('ascii'), 'total_bytes': size,
                'next_offset': offset + len(data) if offset + len(data) < size else None}

    def write(self, root, relative, content, encoding='utf-8', expected_sha256=None):
        with self.lock:
            p = self.mutation_path(root, relative)
            if encoding == 'utf-8':
                data = content.encode('utf-8')
            elif encoding == 'base64':
                if len(content) > LIMIT * 4 // 3 + 8:
                    raise ValueError('File exceeds 10 MB')
                data = base64.b64decode(content, validate=True)
            else:
                raise ValueError('Encoding must be utf-8 or base64')
            if len(data) > LIMIT or (p.exists() and (not p.is_file() or p.stat().st_size > LIMIT)):
                raise ValueError('Writes and overwritten files are limited to 10 MB')
            if not p.parent.is_dir():
                raise ValueError('Create the parent folder first')
            current = self.fingerprint(p)
            # An empty digest explicitly requires an absent destination.
            if expected_sha256 is not None and current != (None if expected_sha256 == '' else expected_sha256):
                raise ValueError('File changed; read again before writing')
            backup = self._backup(p, root, relative, 'write') if p.exists() else None
            self.audit('write', root, relative, status='begin', request_id=backup or '')
            fd, temp = tempfile.mkstemp(prefix='celious-', suffix='.tmp', dir=p.parent)
            try:
                with os.fdopen(fd, 'wb') as out:
                    out.write(data)
                    out.flush()
                    os.fsync(out.fileno())
                self.mutation_path(root, relative)
                if self.fingerprint(p) != current:
                    raise ValueError('File changed during write')
                os.replace(temp, p)
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)
            self.audit('write', root, relative, request_id=backup or '')
            return {'status': 'saved', 'path': relative, 'bytes': len(data),
                    'sha256': hashlib.sha256(data).hexdigest(), 'backup_id': backup}

    def edit(self, root, relative, old_text, new_text, expected_sha256=None):
        with self.lock:
            p = self.mutation_path(root, relative)
            if not p.is_file() or p.stat().st_size > LIMIT:
                raise ValueError('Edit requires a UTF-8 file below 10 MB')
            data = p.read_bytes()
            text = data.decode('utf-8')
            if not old_text or text.count(old_text) != 1:
                raise ValueError('old_text must occur exactly once')
            digest = hashlib.sha256(data).hexdigest()
            if expected_sha256 is not None and digest != expected_sha256:
                raise ValueError('File changed before edit')
            return self.write(root, relative, text.replace(old_text, new_text, 1), expected_sha256=digest)

    def mkdir(self, root, relative):
        with self.lock:
            p = self.mutation_path(root, relative)
            self.audit('mkdir', root, relative, status='begin')
            p.mkdir(parents=True, exist_ok=True)
            self.audit('mkdir', root, relative)
            return {'status': 'created', 'path': relative}

    def _validate_tree(self, root, p):
        # Reject every link/hidden path before moving or deleting a folder.
        for folder, dirs, names in os.walk(p, followlinks=False):
            for name in dirs + names:
                self.path(root, str((Path(folder) / name).relative_to(self.roots[root]['path'])), write=True)

    def move(self, root, relative, destination):
        with self.lock:
            p = self.mutation_path(root, relative)
            dest = self.mutation_path(root, destination)
            if dest.exists() or not dest.parent.is_dir() or not p.exists():
                raise ValueError('Source must exist; destination must be unused with an existing parent')
            if dest.resolve().is_relative_to(p.resolve()):
                raise ValueError('Cannot move a folder into itself')
            if p.is_dir():
                self._validate_tree(root, p)
            self.audit('move', root, relative, status='begin')
            # On Windows rename fails rather than overwriting an existing destination.
            p.rename(dest)
            self.audit('move', root, relative)
            return {'status': 'moved', 'source': relative, 'destination': destination}

    def delete(self, root, relative, recursive=False):
        with self.lock:
            p = self.mutation_path(root, relative)
            if not p.exists():
                raise FileNotFoundError('Target does not exist')
            if p.is_dir():
                self._validate_tree(root, p)
                if not recursive and any(p.iterdir()):
                    raise ValueError('Nonempty folders require recursive=true')
            ident = uuid.uuid4().hex
            folder = self.runtime / 'trash' / ident
            folder.mkdir(parents=True)
            record = {'root': root, 'path': relative, 'status': 'prepared'}
            (folder / 'record.json').write_text(json.dumps(record), 'utf-8')
            self.audit('delete', root, relative, status='begin', request_id=ident)
            # Recoverable deletion: same-volume rename, never recursive filesystem removal.
            self.mutation_path(root, relative)
            p.rename(folder / 'payload')
            record['status'] = 'deleted'
            (folder / 'record.json').write_text(json.dumps(record), 'utf-8')
            self.audit('delete', root, relative, request_id=ident)
            return {'status': 'deleted', 'deletion_id': ident, 'recoverable': True}

    def restore(self, ident):
        with self.lock:
            if not re.fullmatch(r'[a-f0-9]{32}', ident):
                raise ValueError('Invalid deletion ID')
            folder = self.runtime / 'trash' / ident
            record = json.loads((folder / 'record.json').read_text('utf-8'))
            if record['status'] != 'deleted':
                raise ValueError('Deletion was already restored or incomplete')
            p = self.mutation_path(record['root'], record['path'])
            payload = folder / 'payload'
            # Check the saved payload for links as well; runtime can be changed by local software.
            self._no_links(payload)
            for walk_dir, dirs, names in os.walk(payload, followlinks=False):
                for name in dirs + names:
                    self._no_links(Path(walk_dir) / name)
            if p.exists() or not p.parent.is_dir():
                raise ValueError('Restore requires an unused original path and an existing parent')
            self.audit('restore', record['root'], record['path'], status='begin', request_id=ident)
            payload.rename(p)
            record['status'] = 'restored'
            (folder / 'record.json').write_text(json.dumps(record), 'utf-8')
            self.audit('restore', record['root'], record['path'], request_id=ident)
            return {'status': 'restored', 'root': record['root'], 'path': record['path']}

    def list(self, root, path=''):
        p = self.path(root, path)
        result = []
        for item in sorted(p.iterdir(), key=lambda p: p.name.lower()):
            if item.name.startswith('.'):
                continue
            try:
                self.path(root, str(item.relative_to(self.roots[root]['path'])))
                result.append({'name': item.name, 'type': 'directory' if item.is_dir() else 'file'})
            except PermissionError:
                continue
            if len(result) >= 500:
                break
        self.audit('list', root, path)
        return result

    def read(self, root, path, offset=0, max_chars=20000):
        p = self.path(root, path)
        if not p.is_file() or p.stat().st_size > LIMIT:
            raise ValueError('Not a file, or exceeds the 10 MB limit')
        suffix = p.suffix.lower()
        if suffix in TEXT:
            text = p.read_text('utf-8-sig')
            if '\x00' in text:
                raise ValueError('Binary file is not supported as text')
        elif suffix == '.docx':
            from docx import Document
            self._check_zip(p)
            doc = Document(p)
            text = '\n'.join([x.text for x in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif suffix == '.odt':
            self._check_zip(p)
            with zipfile.ZipFile(p) as z:
                tree = ET.fromstring(z.read('content.xml'))
                text = '\n'.join(''.join(e.itertext()) for e in tree.iter() if e.tag.rsplit('}', 1)[-1] in {'p', 'h'})
        elif suffix == '.pdf':
            from pypdf import PdfReader
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(p).pages)
        elif suffix in {'.ods', '.xlsx'}:
            from spreadsheets import load
            sheets = load(p)
            chunks = []
            for sheet in sheets:
                chunks.append('Sheet: ' + sheet['name'])
                rows = {}
                for (r, c), value in sorted(sheet['cells'].items()):
                    rows.setdefault(r, []).append(f'{c}=' + str(value.get('formula') or value['value']))
                for row, values in rows.items():
                    chunks.append(str(row) + ': ' + ' | '.join(values))
            text = '\n'.join(chunks)
        else:
            raise ValueError('Supported: UTF-8 text, Markdown, DOCX, ODT, PDF text, ODS, XLSX. No OCR or legacy Office formats.')
        offset = max(0, offset)
        max_chars = min(50000, max(1, max_chars))
        self.audit('read', root, path)
        return {'text': text[offset:offset + max_chars], 'total_chars': len(text),
                'next_offset': offset + max_chars if offset + max_chars < len(text) else None,
                'note': 'Extracted text; rich-document layout, comments and images may be omitted.' if suffix not in TEXT else ''}

    @staticmethod
    def _check_zip(p):
        with zipfile.ZipFile(p) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Document expanded size exceeds limit')

    def search(self, root, query, path='', content=False):
        if not query or len(query) > 200:
            raise ValueError('Query must have 1 to 200 characters')
        start = self.path(root, path)
        hits, scanned = [], 0
        for folder, dirs, names in os.walk(start, followlinks=False):
            safe = []
            for d in dirs:
                try:
                    self.path(root, str((Path(folder) / d).relative_to(self.roots[root]['path'])))
                    safe.append(d)
                except PermissionError:
                    pass
            dirs[:] = safe
            for name in names:
                scanned += 1
                relative = str((Path(folder) / name).relative_to(self.roots[root]['path']))
                try:
                    p = self.path(root, relative)
                    matched = query.casefold() in name.casefold()
                    if content and not matched:
                        matched = query.casefold() in self.read(root, relative, max_chars=50000)['text'].casefold()
                    if matched:
                        hits.append(relative)
                except (PermissionError, ValueError, OSError, UnicodeError, zipfile.BadZipFile):
                    continue
                if len(hits) >= 100 or scanned >= 5000:
                    self.audit('search', root, path)
                    return {'matches': hits, 'bounded': True, 'scanned': scanned}
        self.audit('search', root, path)
        return {'matches': hits, 'bounded': False, 'scanned': scanned}

    @staticmethod
    def fingerprint(p):
        return hashlib.sha256(p.read_bytes()).hexdigest() if p.is_file() else None

    def stage(self, operation, root, path, text='', destination=''):
        with self.lock:
            p = self.path(root, path, write=True)
            if p == self.roots[root]['path']:
                raise PermissionError('Cannot modify a root directory')
            if operation not in {'write', 'mkdir', 'move'}:
                raise ValueError('Unknown operation; deletion is not supported')
            if operation != 'mkdir' and p.suffix.lower() not in WRITE:
                raise ValueError('Writes and moves are limited to Markdown, TXT, CSV, TSV and JSON')
            if len(text.encode('utf-8')) > 1024 * 1024:
                raise ValueError('Proposed text exceeds 1 MB')
            if p.exists() and (not p.is_file() or p.stat().st_size > 1024 * 1024):
                raise ValueError('Existing target must be a text file below 1 MB')
            if not p.parent.is_dir():
                raise ValueError('Create and approve the parent directory first')
            before = p.read_text('utf-8') if p.exists() else ''
            if operation == 'move':
                dest = self.path(root, destination, write=True)
                if not p.is_file() or dest.exists() or not dest.parent.is_dir() or dest.suffix.lower() not in WRITE:
                    raise ValueError('Move requires an existing text file and unused text destination in an existing folder')
            ident = uuid.uuid4().hex
            request = dict(id=ident, operation=operation, root=root, path=path, text=text,
                           destination=destination, before_hash=self.fingerprint(p), status='pending')
            (self.runtime / f'{ident}.json').write_text(json.dumps(request, indent=2), 'utf-8')
            self.audit('stage_' + operation, root, path, request_id=ident)
            return {'request_id': ident, 'status': 'pending_local_approval',
                    'preview': ''.join(difflib.unified_diff(before.splitlines(True), text.splitlines(True), fromfile=path, tofile=path)) if operation == 'write' else f'{operation}: {path} -> {destination}',
                    'next_step': 'Review locally with service/.venv/Scripts/python.exe service/review.py <request_id>. ChatGPT cannot apply requests.'}

    def request(self, ident):
        if not re.fullmatch(r'[a-f0-9]{32}', ident):
            raise ValueError('Invalid request ID')
        return json.loads((self.runtime / f'{ident}.json').read_text('utf-8'))

    def apply(self, ident):
        # Local CLI only. This method is intentionally NOT an MCP tool.
        with self.lock:
            r = self.request(ident)
            if r['status'] != 'pending':
                raise ValueError('Request was already handled')
            p = self.path(r['root'], r['path'], write=True)
            if self.fingerprint(p) != r['before_hash']:
                raise ValueError('File changed since staging; prepare a fresh request')
            operation = r['operation']
            self.audit('apply_' + operation, r['root'], r['path'], status='begin', request_id=ident)
            if operation == 'write':
                if p.exists():
                    (self.runtime / f'{ident}.backup').write_bytes(p.read_bytes())
                fd, name = tempfile.mkstemp(prefix='celious-', suffix='.tmp', dir=p.parent)
                try:
                    with os.fdopen(fd, 'w', encoding='utf-8', newline='') as out:
                        out.write(r['text'])
                    self.path(r['root'], r['path'], write=True)
                    if self.fingerprint(p) != r['before_hash']:
                        raise ValueError('File changed during approval')
                    os.replace(name, p)
                finally:
                    if os.path.exists(name):
                        os.unlink(name)
            elif operation == 'mkdir':
                p.mkdir(exist_ok=False)
            elif operation == 'move':
                dest = self.path(r['root'], r['destination'], write=True)
                if dest.exists():
                    raise FileExistsError('Destination exists')
                p.rename(dest)
            else:
                raise ValueError('Invalid operation')
            r['status'] = 'applied'
            (self.runtime / f'{ident}.json').write_text(json.dumps(r, indent=2), 'utf-8')
            self.audit('apply_' + operation, r['root'], r['path'], request_id=ident)
            return r['status']
