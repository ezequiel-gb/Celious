import asyncio
import json
import shutil
from pathlib import Path
import sys
import tempfile
import unittest
from engine import Files, BASE

class PolicyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=BASE)
        self.base = Path(self.temp.name)
        self.vault = self.base / 'vault'
        self.source = self.base / 'source'
        self.vault.mkdir()
        self.source.mkdir()
        self.files = Files({'direct_operations': True, 'roots': {'Vault': {'path': str(self.vault), 'writable': True},
                                      'Source': {'path': str(self.source), 'writable': False}}}, self.base / 'runtime')

    def tearDown(self):
        self.temp.cleanup()

    def test_complete_workflow_and_audit(self):
        r = self.files.stage('write', 'Vault', 'test.md', '# Test\nOriginal\n')
        self.assertFalse((self.vault / 'test.md').exists())
        self.files.apply(r['request_id'])
        self.assertIn('Original', self.files.read('Vault', 'test.md')['text'])
        r = self.files.stage('write', 'Vault', 'test.md', '# Test\nModified\n')
        self.files.apply(r['request_id'])
        self.assertIn('Modified', self.files.read('Vault', 'test.md')['text'])
        self.assertIn('Original', (self.files.runtime / (r['request_id'] + '.backup')).read_text())
        r = self.files.stage('mkdir', 'Vault', 'new-folder')
        self.files.apply(r['request_id'])
        r = self.files.stage('move', 'Vault', 'test.md', destination='new-folder/moved.md')
        self.files.apply(r['request_id'])
        self.assertFalse((self.vault / 'test.md').exists())
        self.assertEqual(self.files.search('Vault', 'Modified', content=True)['matches'], ['new-folder\\moved.md'] if sys.platform == 'win32' else ['new-folder/moved.md'])
        self.assertEqual(self.files.list('Vault')[0]['type'], 'directory')
        events = [json.loads(x) for x in (self.files.runtime / 'audit.jsonl').read_text().splitlines()]
        self.assertTrue(any(x['action'] == 'apply_write' for x in events))
        self.assertFalse(any('Modified' in json.dumps(x) for x in events))

    def test_path_boundaries(self):
        for path in ['../outside.txt', '..\\outside.txt', 'C:\\Windows\\win.ini', '/Windows/win.ini', '\\\\localhost\\C$\\Windows', 'a.txt:secret', '.env', 'CON', 'folder./file.txt', 'a/../../b']:
            with self.subTest(path=path), self.assertRaises(PermissionError):
                self.files.path('Vault', path)
        with self.assertRaises(PermissionError):
            self.files.path('Unknown')

    def test_readonly_and_replay(self):
        with self.assertRaises(PermissionError):
            self.files.stage('write', 'Source', 'test.md', 'blocked')
        self.files.config['read_only'] = True
        with self.assertRaises(PermissionError):
            self.files.stage('mkdir', 'Vault', 'blocked')
        self.files.config['read_only'] = False
        r = self.files.stage('write', 'Vault', 'test.md', 'one')
        self.files.config['read_only'] = True
        with self.assertRaises(PermissionError):
            self.files.apply(r['request_id'])
        self.files.config['read_only'] = False
        self.files.apply(r['request_id'])
        with self.assertRaises(ValueError):
            self.files.apply(r['request_id'])

    def test_conflicts_and_unsupported_writes(self):
        p = self.vault / 'test.md'
        p.write_text('original')
        r = self.files.stage('write', 'Vault', 'test.md', 'new')
        p.write_text('human update')
        with self.assertRaises(ValueError):
            self.files.apply(r['request_id'])
        self.assertEqual(p.read_text(), 'human update')
        for extension in ['.ps1', '.exe', '.py', '.docx']:
            with self.assertRaises(ValueError):
                self.files.stage('write', 'Vault', 'test' + extension, 'blocked')
        with self.assertRaises(PermissionError):
            self.files.stage('mkdir', 'Vault', '')

    def test_hardlink_blocked(self):
        import os
        outside = self.source / 'secret.md'
        outside.write_text('secret')
        os.link(outside, self.vault / 'link.md')
        with self.assertRaises(PermissionError):
            self.files.read('Vault', 'link.md')
        for action in [lambda: self.files.write('Vault', 'link.md', 'blocked'),
                       lambda: self.files.delete('Vault', 'link.md'),
                       lambda: self.files.move('Vault', 'link.md', 'moved.md')]:
            with self.assertRaises(PermissionError):
                action()

    def test_direct_operations_any_file_type_and_recovery(self):
        import base64
        self.files.mkdir('Vault', 'folder/nested')
        saved = self.files.write('Vault', 'folder/nested/code.py', 'one\r\ntwo\r\n')
        self.files.edit('Vault', 'folder/nested/code.py', 'two', 'three', saved['sha256'])
        self.assertEqual((self.vault / 'folder/nested/code.py').read_bytes(), b'one\r\nthree\r\n')
        data = b'\x00\xffBinary sample'
        self.files.write('Vault', 'sample.docx', base64.b64encode(data).decode(), 'base64')
        self.assertEqual(base64.b64decode(self.files.raw_read('Vault', 'sample.docx')['base64']), data)
        self.files.move('Vault', 'sample.docx', 'folder/sample.docx')
        with self.assertRaises(ValueError):
            self.files.delete('Vault', 'folder')
        deletion = self.files.delete('Vault', 'folder', recursive=True)
        self.assertFalse((self.vault / 'folder').exists())
        self.files.restore(deletion['deletion_id'])
        self.assertEqual((self.vault / 'folder/sample.docx').read_bytes(), data)
        with self.assertRaises(ValueError):
            self.files.restore(deletion['deletion_id'])
        deletion = self.files.delete('Vault', 'folder/sample.docx')
        self.files.write('Vault', 'folder/sample.docx', 'new file')
        with self.assertRaises(ValueError):
            self.files.restore(deletion['deletion_id'])
        self.assertEqual((self.vault / 'folder/sample.docx').read_text(), 'new file')

    def test_direct_boundaries_and_modes(self):
        for path in ['', '../outside', 'C:\\Windows\\win.ini', '.env']:
            for action in [lambda: self.files.write('Vault', path, 'blocked'),
                           lambda: self.files.mkdir('Vault', path),
                           lambda: self.files.delete('Vault', path, True),
                           lambda: self.files.move('Vault', path, 'elsewhere')]:
                with self.subTest(path=path), self.assertRaises(PermissionError):
                    action()
        with self.assertRaises(PermissionError):
            self.files.write('Source', 'test.txt', 'blocked')
        self.files.config['read_only'] = True
        with self.assertRaises(PermissionError):
            self.files.mkdir('Vault', 'blocked')
        self.files.config['read_only'] = False
        self.files.config['direct_operations'] = False
        with self.assertRaises(PermissionError):
            self.files.write('Vault', 'blocked.txt', 'blocked')

    def test_direct_conflicts_and_backups(self):
        saved = self.files.write('Vault', 'test.txt', 'one')
        updated = self.files.write('Vault', 'test.txt', 'two', expected_sha256=saved['sha256'])
        backup = self.files.runtime / 'backups' / updated['backup_id'] / 'original.bin'
        self.assertEqual(backup.read_bytes(), b'one')
        with self.assertRaises(ValueError):
            self.files.write('Vault', 'test.txt', 'three', expected_sha256=saved['sha256'])
        self.assertEqual((self.vault / 'test.txt').read_text(), 'two')
        self.files.write('Vault', 'other.txt', 'other')
        with self.assertRaises(ValueError):
            self.files.move('Vault', 'test.txt', 'other.txt')
        with self.assertRaises(ValueError):
            self.files.edit('Vault', 'test.txt', 'absent', 'new')

    def test_document_extraction(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        doc = Document()
        doc.add_paragraph('Document sample')
        doc.save(self.vault / 'sample.docx')
        self.assertIn('Document sample', self.files.read('Vault', 'sample.docx')['text'])
        book = Workbook()
        book.active.append(['Sheet sample', '=1+1'])
        book.save(self.vault / 'sample.xlsx')
        self.assertIn('=1+1', self.files.read('Vault', 'sample.xlsx')['text'])
        pdf = PdfWriter()
        pdf.add_blank_page(width=100, height=100)
        pdf.write(self.vault / 'sample.pdf')
        self.assertEqual(self.files.read('Vault', 'sample.pdf')['text'], '')
        import zipfile
        with zipfile.ZipFile(self.vault / 'sample.odt', 'w') as archive:
            archive.writestr('content.xml', '<office:document-content '
                'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
                '<office:body><office:text><text:p>ODT sample</text:p>'
                '</office:text></office:body></office:document-content>')
        self.assertIn('ODT sample', self.files.read('Vault', 'sample.odt')['text'])


async def protocol_test():
    # Installed verification must never inherit operational roots or memory.
    with tempfile.TemporaryDirectory(prefix='celious-file-protocol-') as temporary:
        base = Path(temporary)
        service = base / 'service'
        service.mkdir()
        for name in ['engine.py', 'server.py', 'spreadsheets.py', 'documents.py', 'memory_tools.py', 'project_memory.py', 'second_brain.py',
                     'operations_state.py', 'ingestion.py', 'workflows.py', 'workspace_tools.py', 'authorization.py', 'diagnostics.py', 'connection_settings.py']:
            shutil.copy2(BASE / name, service / name)
        for name in ['Vault', 'Projects']:
            (base / name / 'Inbox').mkdir(parents=True)
        (base / 'Vault/INDEX.md').write_text('# Disposable protocol fixture', 'utf-8')
        config = dict(direct_operations=True, roots={name: dict(path=str(base/name), writable=True) for name in ['Vault', 'Projects']})
        (service / 'config.json').write_text(json.dumps(config), 'utf-8')
        await file_protocol(service, config)


async def file_protocol(service, config):
    from mcp import Client
    from mcp.client.stdio import StdioServerParameters
    transport = StdioServerParameters(command=sys.executable, args=[str(service / 'server.py')])
    async with Client(transport) as client:
        tools = await client.list_tools()
        names = {x.name for x in tools.tools}
        assert names == {'list_roots', 'list_directory', 'read_document', 'search_files', 'propose_write', 'propose_folder', 'propose_move',
                         'read_file_bytes', 'write_file', 'edit_file', 'create_directory', 'move_file', 'delete_file', 'restore_deleted',
                         'inspect_spreadsheet', 'read_spreadsheet', 'search_spreadsheet', 'convert_spreadsheet', 'edit_spreadsheet', 'write_spreadsheet',
                         'workspace_context', 'integration_status', 'workflow_create', 'workflow_run', 'workflow_status', 'workflow_list',
                         'workflow_receipt', 'workflow_reconcile', 'workflow_rollback', 'ingest_preview', 'ingest_commit', 'inspect_archive', 'extract_archive_preview'}, names
        result = await client.call_tool('list_roots', {})
        assert not result.is_error, result
        result = await client.call_tool('read_document', {'root': 'Vault', 'path': 'INDEX.md'})
        assert not result.is_error, result
        result = await client.call_tool('read_document', {'root': 'Vault', 'path': '../service/config.json'})
        assert result.is_error, result
        import uuid
        folder = 'Inbox/connection-test-' + uuid.uuid4().hex
        for tool, args in [
            ('create_directory', {'root':'Vault', 'path':folder}),
            ('write_file', {'root':'Vault', 'path':folder+'/test.py', 'content':'original'}),
            ('edit_file', {'root':'Vault', 'path':folder+'/test.py', 'old_text':'original', 'new_text':'modified'}),
            ('read_document', {'root':'Vault', 'path':folder+'/test.py'}),
            ('move_file', {'root':'Vault', 'path':folder+'/test.py', 'destination':folder+'/renamed.py'}),
            ('delete_file', {'root':'Vault', 'path':folder, 'recursive':True}),
        ]:
            result = await client.call_tool(tool, args)
            assert not result.is_error, (tool, result)
        assert not Files(config, service / '.runtime').path('Vault', folder).exists()
        for tool, args in [('delete_file', {'root':'Vault', 'path':'..', 'recursive':True}),
                           ('write_file', {'root':'Projects', 'path':'../outside.txt', 'content':'blocked'})]:
            result = await client.call_tool(tool, args)
            assert result.is_error, (tool, result)
    print('MCP discovery, direct create/write/edit/read/move/delete and access-denial checks passed; server stopped cleanly.')

if __name__ == '__main__':
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(PolicyTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if not result.wasSuccessful():
        sys.exit(1)
    asyncio.run(protocol_test())
