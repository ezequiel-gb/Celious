"""Synthetic user workflows and refusal cases for native document operations."""
import base64
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from docx import Document
from documents import create, edit, extract, parse, MIME, W
from engine import Files


RECIPE = json.dumps(dict(title='Project handover', blocks=[
    dict(type='heading', text='Continuity', level=1),
    dict(type='paragraph', text='Decision: retain evidence.'),
    dict(type='table', rows=[['Item', 'Status'], ['Source', 'Reviewed']])]))


def repack(data, name, value):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist(): target.writestr(entry, value if entry.filename == name else source.read(entry))
        if name not in source.namelist(): target.writestr(name, value)
    return output.getvalue()


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='celious-documents-')
        self.base = Path(self.temp.name)
        self.root = self.base/'project'
        self.root.mkdir()
        self.files = Files(dict(direct_operations=True, roots=dict(Project=dict(path=str(self.root), writable=True))), self.base/'runtime')

    def tearDown(self): self.temp.cleanup()

    def test_create_read_edit_recover_both_formats(self):
        for suffix in ['.docx', '.odt']:
            with self.subTest(suffix=suffix):
                name = 'handover'+suffix
                saved = self.files.write('Project', name, RECIPE, 'document', expected_sha256='')
                before = (self.root/name).read_bytes()
                self.assertIn('Reviewed', self.files.read('Project', name)['text'])
                changed = self.files.edit('Project', name, 'retain evidence', 'retain attributed evidence', saved['sha256'])
                self.assertTrue(changed['backup_id'])
                self.assertIn('retain attributed evidence', self.files.read('Project', name)['text'])
                backups = list((self.base/'runtime/backups').glob('*/original.bin'))
                self.assertTrue(any(p.read_bytes() == before for p in backups))
                self.files.write('Project', 'recovered'+suffix, base64.b64encode(before).decode(), 'base64', expected_sha256='')
                self.assertEqual((self.root/('recovered'+suffix)).read_bytes(), before)
                self.assertEqual(hashlib.sha256(before).hexdigest(), saved['sha256'])

    def test_split_styled_runs_and_unmodified_parts(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run('retain ').bold = True
        paragraph.add_run('evidence').italic = True
        doc.sections[0].header.paragraphs[0].text = 'Header remains'
        out = io.BytesIO(); doc.save(out)
        before = out.getvalue()
        changed = edit(before, '.docx', 'retain evidence', 'preserve sources')
        with zipfile.ZipFile(io.BytesIO(before)) as a, zipfile.ZipFile(io.BytesIO(changed)) as b:
            for name in a.namelist():
                if name != 'word/document.xml': self.assertEqual(a.read(name), b.read(name), name)
        result = Document(io.BytesIO(changed))
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)
        self.assertEqual(result.paragraphs[0].text, 'preserve sources')
        self.assertIn('Updated header', extract(edit(changed,'.docx','Header remains','Updated header'), '.docx'))

    def test_odt_whitespace_and_span_preservation(self):
        text = '  Café\tcolumns\nnext  line '
        data = create(json.dumps(dict(blocks=[dict(type='paragraph',text=text)])), '.odt')
        self.assertEqual(extract(data,'.odt'), text)
        with zipfile.ZipFile(io.BytesIO(data)) as source:
            self.assertEqual(source.infolist()[0].filename, 'mimetype')
            self.assertEqual(source.read('mimetype').decode(), MIME)
        data=create(RECIPE,'.odt')
        with zipfile.ZipFile(io.BytesIO(data)) as source:
            xml=source.read('content.xml').replace(b'retain evidence', b'<text:span text:style-name="Strong">retain </text:span>evidence')
        changed=edit(repack(data,'content.xml',xml),'.odt','retain evidence','preserve sources')
        self.assertIn('preserve sources', extract(changed,'.odt'))
        with zipfile.ZipFile(io.BytesIO(changed)) as source: self.assertIn(b'Strong',source.read('content.xml'))

    def test_rejected_edits_do_not_change_files(self):
        saved=self.files.write('Project','safe.docx',RECIPE,'document')
        before=(self.root/'safe.docx').read_bytes()
        for old,new,digest in [('missing','new',None),('retain evidence','bad\nline',None),('retain evidence','new','stale')]:
            with self.assertRaises(ValueError): self.files.edit('Project','safe.docx',old,new,digest)
        self.assertEqual((self.root/'safe.docx').read_bytes(),before)
        self.assertFalse(list((self.base/'runtime/backups').glob('*/original.bin')))
        with self.assertRaises(ValueError): edit(create(json.dumps(dict(blocks=[dict(type='paragraph',text='same same')])) ,'.docx'),'.docx','same','different')

    def test_security_and_unsupported_complex_documents(self):
        data=create(RECIPE,'.docx')
        for name in ['word/vbaProject.bin','_xmlsignatures/sig1.xml']:
            with self.assertRaises(ValueError): edit(repack(data,name,b'opaque'),'.docx','retain evidence','new')
        with zipfile.ZipFile(io.BytesIO(data)) as source:
            xml=source.read('word/document.xml').replace(b'<w:p>', b'<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>',1)
        # The heading has a field, but the target paragraph does not: scoped edits remain possible.
        changed=edit(repack(data,'word/document.xml',xml),'.docx','retain evidence','new')
        self.assertIn('new',extract(changed,'.docx'))
        with self.assertRaises(ValueError): edit(repack(data,'word/document.xml',xml),'.docx','Continuity','new')
        with self.assertRaises(ValueError): parse(b'<!DOCTYPE x [<!ENTITY y SYSTEM "file:///private">]><x>&y;</x>')
        with self.assertRaises(ValueError): parse(('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY y "hidden">]><x>&y;</x>').encode('utf-16'))
        odt=create(RECIPE,'.odt')
        with zipfile.ZipFile(io.BytesIO(odt)) as source:
            xml=source.read('content.xml').replace(b'retain evidence',b'retain<!-- private comment --> evidence')
        self.assertNotIn('private comment',extract(repack(odt,'content.xml',xml),'.odt'))
        manifest='<?xml version="1.0" encoding="UTF-16"?><manifest xmlns="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><encryption-data/></manifest>'.encode('utf-16')
        with self.assertRaises(ValueError): edit(repack(odt,'META-INF/manifest.xml',manifest),'.odt','retain evidence','new')
        with self.assertRaises(ValueError): edit(repack(odt,'Basic/Standard/module.xml',b'opaque'),'.odt','retain evidence','new')
        with zipfile.ZipFile(io.BytesIO(odt)) as source:
            xml=source.read('content.xml').replace(b'<office:body>',b'<office:scripts/><office:body>')
        with self.assertRaises(ValueError): edit(repack(odt,'content.xml',xml),'.odt','retain evidence','new')
        with self.assertRaises(PermissionError): self.files.write('Project','../outside.docx',RECIPE,'document')
        self.files.config['read_only']=True
        with self.assertRaises(PermissionError): self.files.write('Project','blocked.odt',RECIPE,'document')

    def test_invalid_recipes(self):
        for recipe in ['{}','[]','{"blocks":[]}',json.dumps(dict(blocks=[dict(type='table',rows=[['a'],['a','b']])])),json.dumps(dict(blocks=[dict(type='heading',text='bad',level=7)]))]:
            with self.subTest(recipe=recipe), self.assertRaises(ValueError): create(recipe,'.docx')
        with self.assertRaises(ValueError): create(RECIPE,'.pdf')

    def test_native_workflow_crash_resume_and_reverse_recovery(self):
        from workflows import Workflows
        for suffix in ['.docx', '.odt']:
            with self.subTest(suffix=suffix):
                flow=Workflows(self.files)
                name='resumable'+suffix
                plan=flow.create('Document recovery',[
                    dict(id='save',op='write',args=dict(root='Project',relative=name,content=RECIPE,encoding='document'),needs=[]),
                    dict(id='edit',op='edit',args=dict(root='Project',relative=name,old_text='retain evidence',new_text='retain exact evidence'),needs=['save'])])
                flow._interrupt_after_effect=True
                with self.assertRaises(SystemExit): flow.run(plan['id'],plan['plan_sha256'])
                initial=(self.root/name).read_bytes()
                flow=Workflows(self.files)
                flow._interrupt_after_effect=True
                with self.assertRaises(SystemExit): flow.run(plan['id'],plan['plan_sha256'])
                edited=(self.root/name).read_bytes()
                self.assertNotEqual(initial,edited)
                flow=Workflows(self.files)
                self.assertEqual(flow.run(plan['id'],plan['plan_sha256'])['status'],'completed')
                self.assertEqual((self.root/name).read_bytes(),edited)
                self.assertIn('retain exact evidence',self.files.read('Project',name)['text'])
                self.assertEqual(flow.rollback(plan['id'],plan['plan_sha256'])['status'],'rolled_back')
                self.assertFalse((self.root/name).exists())


if __name__ == '__main__': unittest.main()
