"""Bounded, typed spreadsheet I/O. No formula evaluation or macro execution."""
import base64
import csv
from datetime import date, datetime, time, timedelta
import io
import json
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

FORMATS = {'.ods', '.xlsx', '.csv', '.tsv'}
MAX_CELLS = 200000
MAX_ROWS = 1048576
MAX_COLS = 16384
MAX_TEXT = 5000000
NS = {'o': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
      't': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
      'x': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
LOSS = ('Data transfer only: formatting, charts, macros, comments, validation, links, '
        'names and workbook metadata are not transferred. Formulas require explicit cached-values mode; '
        'cached results can be absent or stale. No recalculation is performed.')


class SpreadsheetError(ValueError):
    """Safe, actionable message without source cell contents."""


def attr(e, ns, key, default=None):
    return e.get('{'+NS[ns]+'}'+key, default)


def scalar(value, formula=None, cached=None):
    if isinstance(value, (datetime, date, time)):
        kind, value = 'date', value.isoformat()
    elif isinstance(value, timedelta):
        kind, value = 'duration', value.total_seconds()
    elif value is None:
        kind = 'empty'
    elif isinstance(value, bool):
        kind = 'boolean'
    elif isinstance(value, (float, int)):
        import math
        if not math.isfinite(value):
            raise SpreadsheetError('Non-finite numeric value is unsupported')
        kind = 'number'
    else:
        kind, value = 'string', str(value)
    result = {'type': kind, 'value': value}
    if formula:
        result.update(formula=formula, cached=cached)
    return result


def sheet(name):
    if not isinstance(name,str) or not name or len(name)>255 or re.search(r'[\x00-\x1f]',name):
        raise SpreadsheetError('Worksheet names must be nonempty text of at most 255 characters without controls')
    return {'name': name, 'rows': 0, 'columns': 0, 'cells': {}, 'merged_ranges': []}


def put(s, r, c, value):
    if not (1 <= r <= MAX_ROWS and 1 <= c <= MAX_COLS):
        raise SpreadsheetError('Worksheet coordinates exceed supported limits')
    if value['value'] is not None or value.get('formula'):
        if len(s['cells']) >= MAX_CELLS:
            raise SpreadsheetError('Workbook exceeds 200000 populated cells; split the source workbook')
        s['cells'][r, c] = value
        s['rows'], s['columns'] = max(s['rows'], r), max(s['columns'], c)


def xml(data):
    # Reject DTDs and entities, including UTF-16/32 encodings, before ElementTree.
    if b'<!DOCTYPE' in data.upper().replace(b'\x00', b'') or b'<!ENTITY' in data.upper().replace(b'\x00', b''):
        raise SpreadsheetError('XML DTDs and entities are prohibited')
    return ET.fromstring(data)


def zipcheck(p):
    with zipfile.ZipFile(p) as z:
        entries = z.infolist()
        if len(entries) > 10000 or sum(i.file_size for i in entries) > 50*1024*1024:
            raise SpreadsheetError('Expanded archive exceeds 50 MB or 10000 entries')
        if len({i.filename for i in entries}) != len(entries) or any(i.flag_bits & 1 for i in entries):
            raise SpreadsheetError('Encrypted or duplicate archive entries are unsupported')
        for i in entries:
            if i.filename.lower().endswith(('.xml', '.rels')):
                data = z.read(i)
                if b'<!DOCTYPE' in data.upper().replace(b'\x00', b'') or b'<!ENTITY' in data.upper().replace(b'\x00', b''):
                    raise SpreadsheetError('XML DTDs and entities are prohibited')


def load(p):
    p = Path(p)
    try:
        if p.suffix.lower() not in FORMATS:
            raise SpreadsheetError('Supported spreadsheets: ODS, XLSX, CSV, TSV. Legacy XLS is unsupported.')
        if not p.is_file() or p.stat().st_size > 10*1024*1024:
            raise SpreadsheetError('Spreadsheet must be a file of at most 10 MB')
        ext = p.suffix.lower()
        sheets = []
        if ext in {'.csv', '.tsv'}:
            text = p.read_text('utf-8-sig')
            if '\x00' in text:
                raise SpreadsheetError('CSV/TSV must contain UTF-8 text, without NUL bytes')
            s = sheet('Sheet1')
            for r, row in enumerate(csv.reader(io.StringIO(text, newline=''), delimiter='\t' if ext == '.tsv' else ',', strict=True), 1):
                if r > MAX_ROWS or len(row) > MAX_COLS:
                    raise SpreadsheetError('Delimited file exceeds worksheet limits')
                for c, v in enumerate(row, 1):
                    if v != '':
                        put(s, r, c, scalar(v))
            sheets.append(s)
        else:
            zipcheck(p)
            if ext == '.ods':
                with zipfile.ZipFile(p) as z:
                    if z.read('mimetype') != b'application/vnd.oasis.opendocument.spreadsheet':
                        raise SpreadsheetError('File is not an OpenDocument spreadsheet')
                    doc = xml(z.read('content.xml'))
                tables = doc.findall('./o:body/o:spreadsheet/t:table', NS)
                for table in tables:
                    s = sheet(attr(table, 't', 'name', ''))
                    r = 1
                    for row in table.iter('{'+NS['t']+'}table-row'):
                        repeat = int(attr(row, 't', 'number-rows-repeated', '1'))
                        if repeat < 1 or r + repeat - 1 > MAX_ROWS:
                            raise SpreadsheetError('Invalid ODS row repetition')
                        c, pending = 1, []
                        for e in row:
                            if e.tag not in {'{'+NS['t']+'}table-cell', '{'+NS['t']+'}covered-table-cell'}:
                                continue
                            n = int(attr(e, 't', 'number-columns-repeated', '1'))
                            if n < 1 or c+n-1 > MAX_COLS:
                                raise SpreadsheetError('Invalid ODS column repetition')
                            typ = attr(e, 'o', 'value-type')
                            if typ not in {None,'string','float','currency','percentage','boolean','date','time'}:
                                raise SpreadsheetError('Unsupported ODS cell type')
                            paragraphs = [''.join(ods_text(t)) for t in e.findall('x:p', NS)]
                            v = '\n'.join(paragraphs) if paragraphs else attr(e, 'o', 'string-value')
                            if typ in {'float', 'currency', 'percentage'}:
                                v = float(attr(e, 'o', 'value'))
                            elif typ == 'boolean':
                                raw = attr(e, 'o', 'boolean-value')
                                if raw not in {'true', 'false'}:
                                    raise SpreadsheetError('Invalid ODS boolean')
                                v = raw == 'true'
                            elif typ == 'date':
                                v = attr(e, 'o', 'date-value')
                                datetime.fromisoformat(v.replace('Z', '+00:00'))
                            elif typ == 'time':
                                v = attr(e, 'o', 'time-value')
                            f = attr(e, 't', 'formula')
                            cell = scalar(v, f, v if f else None)
                            if typ in {'date', 'time'}:
                                cell['type'] = 'date' if typ == 'date' else 'duration'
                            if v is not None or f:
                                if len(pending) + n > MAX_CELLS or repeat*n > MAX_CELLS:
                                    raise SpreadsheetError('ODS repeated data exceeds populated-cell limit')
                                pending.extend((c+i, cell.copy()) for i in range(n))
                            cs, rs = int(attr(e, 't', 'number-columns-spanned', '1')), int(attr(e, 't', 'number-rows-spanned', '1'))
                            if cs < 1 or rs < 1 or c+cs-1 > MAX_COLS or r+repeat+rs-2 > MAX_ROWS:
                                raise SpreadsheetError('Invalid merged-cell span')
                            if cs > 1 or rs > 1:
                                if repeat > 10000 or len(s['merged_ranges']) + repeat > 10000:
                                    raise SpreadsheetError('Too many merged ranges')
                                for rr in range(r, r+repeat):
                                    s['merged_ranges'].append(f'{get_column_letter(c)}{rr}:{get_column_letter(c+cs-1)}{rr+rs-1}')
                            c += n
                        if pending:
                            if len(s['cells']) + len(pending)*repeat > MAX_CELLS:
                                raise SpreadsheetError('ODS data exceeds populated-cell limit')
                            for rr in range(r, r+repeat):
                                for cc, value in pending:
                                    put(s, rr, cc, value.copy())
                        r += repeat
                    sheets.append(s)
            else:
                with zipfile.ZipFile(p) as z:
                    if 'xl/workbook.xml' not in z.namelist():
                        raise SpreadsheetError('File is not an XLSX workbook')
                    stored_cells=0
                    for entry in z.namelist():
                        if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',entry):
                            for _,element in ET.iterparse(io.BytesIO(z.read(entry)),events=('end',)):
                                if element.tag.rsplit('}',1)[-1]=='c':
                                    stored_cells+=1
                                    if stored_cells>MAX_CELLS:
                                        raise SpreadsheetError('XLSX exceeds 200000 stored cells')
                                element.clear()
                book = load_workbook(p, data_only=False, keep_links=False)
                cache = load_workbook(p, data_only=True, keep_links=False)
                try:
                    for ws in book:
                        s = sheet(ws.title)
                        # Sparse XML avoids iterating malicious/inflated dimensions.
                        if len(ws._cells) > MAX_CELLS:
                            raise SpreadsheetError('XLSX exceeds 200000 stored cells')
                        for (r, c), e in ws._cells.items():
                            f = e.value if e.data_type == 'f' else None
                            v = cache[ws.title].cell(r, c).value if f else e.value
                            cell = scalar(v, f, scalar(v)['value'] if f else None)
                            if e.data_type == 'e':
                                cell['type'] = 'error'
                            put(s, r, c, cell)
                        s['merged_ranges'] = [str(x) for x in ws.merged_cells.ranges]
                        sheets.append(s)
                finally:
                    book.close()
                    cache.close()
        if not sheets or len(sheets) > 256 or len({s['name'] for s in sheets}) != len(sheets):
            raise SpreadsheetError('Workbook requires 1 to 256 uniquely named worksheets')
        if sum(len(s['cells']) for s in sheets) > MAX_CELLS:
            raise SpreadsheetError('Workbook exceeds 200000 populated cells')
        if sum(len(str(v['value'])) + len(v.get('formula', '')) for s in sheets for v in s['cells'].values()) > MAX_TEXT:
            raise SpreadsheetError('Workbook text exceeds 5 million characters')
        return sheets
    except SpreadsheetError:
        raise
    except (ValueError, KeyError, TypeError, OverflowError, ET.ParseError, zipfile.BadZipFile, csv.Error, UnicodeError, IndexError):
        raise SpreadsheetError('Malformed spreadsheet, invalid typed value, or mismatched file content') from None


def ods_text(e):
    if e.text:
        yield e.text
    for child in e:
        if child.tag == '{'+NS['x']+'}s':
            count = int(attr(child, 'x', 'c', '1'))
            if count < 1 or count > 10000:
                raise SpreadsheetError('Invalid ODS whitespace repetition')
            yield ' '*count
        elif child.tag == '{'+NS['x']+'}tab':
            yield '\t'
        elif child.tag == '{'+NS['x']+'}line-break':
            yield '\n'
        else:
            yield from ods_text(child)
        if child.tail:
            yield child.tail


def select(sheets, name):
    if name is None:
        if len(sheets) != 1:
            raise SpreadsheetError('Select a worksheet explicitly for a multi-sheet workbook')
        return sheets[0]
    for s in sheets:
        if s['name'] == name:
            return s
    raise SpreadsheetError('Worksheet not found; inspect worksheet names first')


def read(sheets, name=None, start_row=1, row_count=100, start_column=1, column_count=30):
    s = select(sheets, name)
    if min(start_row, start_column, row_count, column_count) < 1 or row_count > 500 or column_count > 100 or row_count*column_count > 10000:
        raise SpreadsheetError('Use positive coordinates, at most 500 rows, 100 columns and 10000 cells per page')
    end_r, end_c = min(s['rows'], start_row+row_count-1), min(s['columns'], start_column+column_count-1)
    rows = []
    size = 0
    for r in range(start_row, end_r+1):
        cells = [dict(address=f'{get_column_letter(c)}{r}', **s['cells'].get((r,c), scalar(None))) for c in range(start_column, end_c+1)]
        size += len(json.dumps(cells, ensure_ascii=False))
        if size > 200000:
            raise SpreadsheetError('Page exceeds 200000 characters; request fewer rows or columns')
        rows.append({'row': r, 'cells': cells})
    return {'worksheet': s['name'], 'rows': rows, 'dimensions': [s['rows'], s['columns']],
            'next_row': end_r+1 if end_r < s['rows'] else None,
            'next_column': end_c+1 if end_c < s['columns'] else None,
            'bounds': 'Dimensions cover populated/formula cells; styled empty tails are excluded.'}


def inspect(sheets):
    result = []
    for s in sheets:
        rs = sorted({r for r,c in s['cells']})
        cs = sorted({c for r,c in s['cells']})
        header = rs[0] if rs else None
        headers = [{'column': c, **s['cells'][header,c]} for c in cs if (header,c) in s['cells']] if header else []
        result.append({'name': s['name'], 'rows': s['rows'], 'columns': s['columns'],
                       'populated_rows': len(rs), 'data_records': max(0, len(rs)-1), 'header_row': header,
                       'headers': headers[:100], 'headers_next_column': headers[100]['column'] if len(headers)>100 else None,
                       'populated_cells': len(s['cells']), 'blank_rows_within_extent': s['rows']-len(rs),
                       'blank_columns_within_extent': s['columns']-len(cs),
                       'formulas': sum(bool(v.get('formula')) for v in s['cells'].values()),
                       'merged_ranges': s['merged_ranges'][:100], 'merged_range_count': len(s['merged_ranges'])})
    if len(json.dumps(result,ensure_ascii=False)) > 200000:
        raise SpreadsheetError('Metadata exceeds response limit; use read_spreadsheet with narrow coordinates')
    return {'worksheets': result, 'record_definition': 'Nonempty rows after first populated row, treated as header; verify header assumption.',
            'limits': {'populated_cells': MAX_CELLS, 'input_bytes': 10485760}, 'formula_evaluation': False}


def search(sheets, name=None, query='', columns=None, mode='contains', offset=0, limit=100):
    s = select(sheets, name)
    if mode not in {'contains', 'duplicates', 'missing', 'inconsistent'} or not 1 <= limit <= 500 or offset < 0 or len(query)>200:
        raise SpreadsheetError('Invalid search mode/pagination/query')
    cols = columns or sorted({c for r,c in s['cells']})
    if not cols or len(cols)>100 or any(not isinstance(c,int) or c<1 or c>MAX_COLS for c in cols):
        raise SpreadsheetError('Select 1 to 100 valid column numbers')
    rows = sorted({r for r,c in s['cells']})
    data = rows[1:]
    def value(r,c):
        return s['cells'].get((r,c), scalar(None))['value']
    hits = []
    if mode == 'duplicates':
        groups = {}
        for r in data:
            key = tuple((s['cells'].get((r,c),scalar(None))['type'], str(value(r,c)).strip().casefold()) for c in cols)
            if all(value(r,c) in (None,'') for c in cols):
                continue
            groups.setdefault(key,[]).append(r)
        hits = [{'rows': rr} for rr in groups.values() if len(rr)>1]
    else:
        expected = {}
        if mode == 'inconsistent':
            from collections import Counter
            for c in cols:
                types = Counter(s['cells'][r,c]['type'] for r in data if (r,c) in s['cells'])
                expected[c] = types.most_common(1)[0][0] if types else None
        if mode == 'contains' and not query:
            raise SpreadsheetError('Contains search requires a nonempty query')
        for r in (rows if mode == 'contains' else data):
            matched = [c for c in cols if (query.casefold() in str(value(r,c) if value(r,c) is not None else '').casefold() if mode == 'contains'
                       else value(r,c) in (None,'') if mode == 'missing'
                       else (r,c) in s['cells'] and s['cells'][r,c]['type'] != expected[c])]
            if matched:
                hits.append({'row': r, 'columns': matched})
    page=hits[offset:offset+limit]
    if len(json.dumps(page))>200000:
        raise SpreadsheetError('Search page too large; reduce limit or narrow columns')
    return {'worksheet': s['name'], 'mode': mode, 'matches': page, 'total': len(hits),
            'next_offset': offset+limit if offset+limit<len(hits) else None,
            'note': 'Duplicate matching trims/casefolds typed values and excludes fully blank keys. Missing/type checks apply to populated data rows.'}


def safe_string(v):
    return isinstance(v,str) and v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n'))


def encode(sheets, ext, worksheet=None, formula_mode='reject'):
    if ext not in FORMATS:
        raise SpreadsheetError('Output format must be ODS, XLSX, CSV or TSV')
    if formula_mode not in {'reject','cached_values'}:
        raise SpreadsheetError('formula_mode must be reject or cached_values')
    if any(v.get('formula') for s in sheets for v in s['cells'].values()) and formula_mode != 'cached_values':
        raise SpreadsheetError('Formula translation is not supported; explicitly choose cached_values to discard formulas')
    if ext in {'.csv','.tsv'}:
        s=select(sheets,worksheet)
        if s['rows']*s['columns']>MAX_CELLS:
            raise SpreadsheetError('Delimited export rectangle exceeds 200000 cells')
        out=io.StringIO(newline='')
        writer=csv.writer(out,delimiter='\t' if ext=='.tsv' else ',')
        for r in range(1,s['rows']+1):
            row=[s['cells'].get((r,c),scalar(None))['value'] for c in range(1,s['columns']+1)]
            if any(safe_string(v) for v in row):
                raise SpreadsheetError('CSV/TSV formula injection risk; use XLSX literal-string output instead')
            writer.writerow(['' if v is None else v for v in row])
        return out.getvalue().encode('utf-8-sig')
    if ext=='.xlsx':
        book=Workbook()
        book.remove(book.active)
        try:
            for s in sheets:
                if len(s['name'])>31 or re.search(r'[\\*?:/\[\]]',s['name']):
                    raise SpreadsheetError('Worksheet name is incompatible with XLSX')
                ws=book.create_sheet(s['name'])
                if ws.title != s['name'] or len(s['name'])>31:
                    raise SpreadsheetError('Worksheet name is incompatible with XLSX')
                for (r,c),v in s['cells'].items():
                    if v['type'] in {'duration','error'}:
                        raise SpreadsheetError('XLSX data export of duration/error cells is unsupported')
                    value=v['value']
                    if v['type']=='date' and value is not None:
                        value=datetime.fromisoformat(value.replace('Z','+00:00'))
                        if value.tzinfo:
                            raise SpreadsheetError('Timezone dates cannot be exported to XLSX safely')
                    cell=ws.cell(r,c,value)
                    if isinstance(value,str):
                        if len(value)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value):
                            raise SpreadsheetError('Text exceeds XLSX cell length or contains prohibited control characters')
                        cell.data_type='s'  # Never introduce a formula from untrusted text.
                for merge in s['merged_ranges']:
                    a,b,c,d=range_boundaries(merge)
                    if (c-a+1)*(d-b+1)>MAX_CELLS:
                        raise SpreadsheetError('Merged range exceeds conversion limit')
                    if any((rr,cc) in s['cells'] for rr in range(b,d+1) for cc in range(a,c+1) if (rr,cc)!=(b,a)):
                        raise SpreadsheetError('Merged range would discard populated cells')
                    ws.merge_cells(merge)
            out=io.BytesIO()
            book.save(out)
            return out.getvalue()
        finally:
            book.close()
    # Minimal standards-compliant ODS data export. No original package parts copied.
    root=ET.Element('{'+NS['o']+'}document-content',{'{'+NS['o']+'}version':'1.2'})
    body=ET.SubElement(root,'{'+NS['o']+'}body')
    ss=ET.SubElement(body,'{'+NS['o']+'}spreadsheet')
    for s in sheets:
        if s['rows']*s['columns']>MAX_CELLS or s['merged_ranges']:
            raise SpreadsheetError('ODS data export requires at most 200000 rectangular cells and no merges')
        table=ET.SubElement(ss,'{'+NS['t']+'}table',{'{'+NS['t']+'}name':s['name']})
        for r in range(1,s['rows']+1):
            row=ET.SubElement(table,'{'+NS['t']+'}table-row')
            for c in range(1,s['columns']+1):
                e=ET.SubElement(row,'{'+NS['t']+'}table-cell')
                v=s['cells'].get((r,c),scalar(None))
                value=v['value']
                if value is None:
                    continue
                typ=v['type']
                if typ=='boolean':
                    e.set('{'+NS['o']+'}value-type','boolean'); e.set('{'+NS['o']+'}boolean-value',str(value).lower())
                elif typ=='number':
                    e.set('{'+NS['o']+'}value-type','float'); e.set('{'+NS['o']+'}value',str(value))
                elif typ=='date':
                    e.set('{'+NS['o']+'}value-type','date'); e.set('{'+NS['o']+'}date-value',value)
                elif typ in {'duration','error'}:
                    raise SpreadsheetError('ODS export of duration/error cells is unsupported')
                else:
                    e.set('{'+NS['o']+'}value-type','string')
                ET.SubElement(e,'{'+NS['x']+'}p').text=str(value)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet',compress_type=zipfile.ZIP_STORED)
        z.writestr('content.xml',ET.tostring(root,encoding='utf-8',xml_declaration=True))
        z.writestr('META-INF/manifest.xml','<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    return out.getvalue()


class Spreadsheets:
    def __init__(self, files):
        self.files=files

    def load(self, root, path):
        sheets=load(self.files.path(root,path))
        self.files.audit('spreadsheet_read',root,path)
        return sheets

    def inspect(self, root, path):
        result=inspect(self.load(root,path))
        p=self.files.path(root,path)
        result.update(sha256=self.files.fingerprint(p),bytes=p.stat().st_size,format=p.suffix.lower())
        return result

    def save(self, root, destination, data, expected_sha256=''):
        import os
        import tempfile
        p=self.files.mutation_path(root,destination)
        if p.suffix.lower() not in FORMATS:
            raise SpreadsheetError('Unsupported spreadsheet destination')
        if expected_sha256 is None:
            raise SpreadsheetError('Use empty digest for a new file or current SHA256 for replacement')
        # Verify serialized output before committing. Temp file is inside existing destination parent.
        fd,tmp=tempfile.mkstemp(prefix='celious-verify-',suffix=p.suffix,dir=p.parent)
        try:
            with os.fdopen(fd,'wb') as out:
                out.write(data); out.flush(); os.fsync(out.fileno())
            verified=load(Path(tmp))
            result=self.files.write(root,destination,base64.b64encode(data).decode('ascii'),'base64',expected_sha256)
            if self.files.fingerprint(p)!=result['sha256']:
                raise SpreadsheetError('Post-write checksum mismatch; recover using returned backup record')
            result.update(verified=True,worksheets=[s['name'] for s in verified])
            return result
        finally:
            Path(tmp).unlink(missing_ok=True)

    def convert(self,root,path,destination,worksheet=None,acknowledge_loss=False,formula_mode='reject'):
        if not acknowledge_loss:
            raise SpreadsheetError('Set acknowledge_loss=true after reviewing: '+LOSS)
        src=self.files.path(root,path)
        dest=self.files.mutation_path(root,destination)
        if src.resolve()==dest.resolve():
            raise SpreadsheetError('Conversion requires a different destination')
        sheets=self.load(root,path)
        data=encode(sheets,dest.suffix.lower(),worksheet,formula_mode)
        result=self.save(root,destination,data)
        result['limitations']=LOSS
        return result

    def write(self,root,destination,workbook_json,allow_formulas=False,expected_sha256=''):
        if len(workbook_json)>1000000:
            raise SpreadsheetError('Workbook request exceeds 1 million characters')
        try:
            spec=json.loads(workbook_json)
            if not isinstance(spec,list) or not 1<=len(spec)<=256:
                raise ValueError()
            sheets=[]
            for item in spec:
                s=sheet(item['name'])
                for r,row in enumerate(item.get('rows',[]),1):
                    for c,v in enumerate(row,1):
                        if isinstance(v,(dict,list)):
                            raise SpreadsheetError('Rows accept JSON scalar values; use edit_spreadsheet for explicit formulas')
                        put(s,r,c,scalar(v))
                sheets.append(s)
            if len({s['name'] for s in sheets})!=len(sheets):
                raise ValueError()
        except (ValueError,KeyError,TypeError):
            raise SpreadsheetError('workbook_json must be a list of unique {name, rows} worksheets containing scalar cells') from None
        return self.save(root,destination,encode(sheets,Path(destination).suffix.lower()),expected_sha256)

    def edit(self,root,path,destination,operations_json,acknowledge_loss=False,allow_formulas=False,expected_sha256=''):
        # Conservative native XLSX editing retains common styles/formulas; acknowledge uncommon-feature risk.
        if not acknowledge_loss:
            raise SpreadsheetError('Editing requires acknowledge_loss=true: XLSX common styles/formulas retained, unsupported features may be lost. ODS/CSV/TSV editing uses data-only export. '+LOSS)
        if len(operations_json)>1000000:
            raise SpreadsheetError('Edit request exceeds 1 million characters')
        src=self.files.path(root,path)
        dest=self.files.mutation_path(root,destination)
        if src.resolve()==dest.resolve() and not expected_sha256:
            raise SpreadsheetError('Replacing source requires its current SHA256; prefer a new destination')
        sheets=self.load(root,path)
        try:
            ops=json.loads(operations_json)
            if not isinstance(ops,list) or not 1<=len(ops)<=1000:
                raise ValueError()
            native=src.suffix.lower()==dest.suffix.lower()=='.xlsx'
            book=load_workbook(src,keep_links=True) if native else None
            try:
                for op in ops:
                    action=op['action']
                    if action=='add_sheet':
                        if any(s['name']==op['sheet'] for s in sheets):
                            raise SpreadsheetError('Worksheet already exists')
                        sheets.append(sheet(op['sheet']))
                        if book: book.create_sheet(op['sheet'])
                        continue
                    s=select(sheets,op['sheet'])
                    if action=='append_rows':
                        changes=[(s['rows']+1+i,c+1,v) for i,row in enumerate(op['rows']) for c,v in enumerate(row)]
                    elif action=='set_cell':
                        changes=[(op['row'],op['column'],op.get('value'))]
                    else:
                        raise SpreadsheetError('Actions: set_cell, append_rows, add_sheet. Set header/column cells with set_cell.')
                    for r,c,v in changes:
                        if not isinstance(r,int) or not isinstance(c,int) or not 1<=r<=MAX_ROWS or not 1<=c<=MAX_COLS or isinstance(v,(dict,list)):
                            raise SpreadsheetError('Invalid cell coordinates or value')
                        f=op.get('formula') if action=='set_cell' else None
                        if f:
                            functions = re.findall(r'([A-Za-z_][A-Za-z_0-9.]*)\s*\(', f) if isinstance(f,str) else []
                            allowed = {'SUM','AVERAGE','MIN','MAX','COUNT','COUNTA','IF','AND','OR','NOT','ROUND','ABS','COUNTIF','SUMIF'}
                            if not book or not allow_formulas or not isinstance(f,str) or not re.fullmatch(r'=[0-9A-Za-z_ :.,()+*/^$%<>=!"\-]+',f) or any(fn.upper() not in allowed for fn in functions) or re.search(r'(?i)https?|file|\\|\[|WEBSERVICE|HYPERLINK|DDE|RTD|CALL|EXEC|REGISTER',f):
                                raise SpreadsheetError('Formula requires native XLSX, allow_formulas=true and a local arithmetic formula without external references')
                        if v is None and not f:
                            s['cells'].pop((r,c),None)
                        else:
                            put(s,r,c,scalar(v,f))
                        if book:
                            e=book[s['name']].cell(r,c)
                            if isinstance(v,str) and (len(v)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',v)):
                                raise SpreadsheetError('Text exceeds XLSX cell length or contains prohibited control characters')
                            e.value=f or v
                            if isinstance(v,str) and not f: e.data_type='s'
                if book:
                    out=io.BytesIO(); book.save(out); data=out.getvalue()
                else:
                    data=encode(sheets,dest.suffix.lower())
                result=self.save(root,destination,data,expected_sha256)
                result['limitations']=LOSS if not native else 'Common XLSX features retained by openpyxl; unsupported extensions may be lost; formulas are not recalculated.'
                return result
            finally:
                if book: book.close()
        except (KeyError,TypeError,json.JSONDecodeError):
            raise SpreadsheetError('Invalid operations_json; see SPREADSHEETS.md') from None


def register(mcp, files, invoke, read_annotation, mutation):
    service=Spreadsheets(files)

    @mcp.tool(annotations=read_annotation)
    def inspect_spreadsheet(root: str,path: str) -> dict:
        """Inspect ODS/XLSX/CSV/TSV worksheet names, populated dimensions, headers, merges and types. No macros or formula evaluation."""
        return invoke(service.inspect,root,path)

    @mcp.tool(annotations=read_annotation)
    def read_spreadsheet(root: str,path: str,worksheet: str | None=None,start_row: int=1,row_count: int=100,start_column: int=1,column_count: int=30) -> dict:
        """Read typed cells/formulas/caches. One-based bounds, <=500 rows/100 columns/10000 cells. Follow next_row/next_column; narrow queries protect private data."""
        return invoke(lambda: read(service.load(root,path),worksheet,start_row,row_count,start_column,column_count))

    @mcp.tool(annotations=read_annotation)
    def search_spreadsheet(root: str,path: str,worksheet: str | None=None,query: str='',columns: list[int] | None=None,mode: str='contains',offset: int=0,limit: int=100) -> dict:
        """Paginated contains, duplicates, missing or inconsistent-type checks. Returns row/column coordinates, not private cell contents. First populated row is header."""
        return invoke(lambda: search(service.load(root,path),worksheet,query,columns,mode,offset,limit))

    @mcp.tool(annotations=mutation)
    def convert_spreadsheet(root: str,path: str,destination: str,worksheet: str | None=None,acknowledge_loss: bool=False,formula_mode: str='reject') -> dict:
        """Convert ODS/XLSX/CSV/TSV to a NEW file. Explicit worksheet for CSV/TSV. Requires loss acknowledgement; formatting/objects discarded. Formulas reject unless cached_values explicitly selected."""
        return invoke(service.convert,root,path,destination,worksheet,acknowledge_loss,formula_mode)

    @mcp.tool(annotations=mutation)
    def write_spreadsheet(root: str,destination: str,workbook_json: str,expected_sha256: str='') -> dict:
        """Create data workbook from JSON [{name,rows:[[scalar,...],...]}]. Text stays literal. Empty digest requires absent destination; replacement needs current SHA256 and creates backup."""
        return invoke(service.write,root,destination,workbook_json,False,expected_sha256)

    @mcp.tool(annotations=mutation)
    def edit_spreadsheet(root: str,path: str,destination: str,operations_json: str,acknowledge_loss: bool=False,allow_formulas: bool=False,expected_sha256: str='') -> dict:
        """Controlled set_cell/append_rows/add_sheet edits. Prefer new destination; replacement requires current SHA256 and backup. Explicit feature-loss acknowledgement required. See SPREADSHEETS.md."""
        return invoke(service.edit,root,path,destination,operations_json,acknowledge_loss,allow_formulas,expected_sha256)
